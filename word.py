import docx
import xlrd
import xlwt
import re
from xlutils.copy import copy
import time

print('开始生成excel:',time.time())

# 必填选填数据
def select_choice(file_name):

    wb = xlrd.open_workbook(file_name)
    tables = wb.sheets()
    li =[]
    for table in tables[1:]:
        nrows = table.nrows #行数
        ncols = table.ncols #列数
        # print(nrows, ncols)
        colnames =  table.row_values(2) #某一行数据
        d = {}
        for rownum in range(4,nrows-3):

            row = table.row_values(rownum)
            if row:
                app = {}
                app['table'] = table.name
                for i in range(len(colnames)-4):
                    app[colnames[i]] = row[i]
            li.append(app)

    return li

# 将world接口文档转为excel
def world_2_excel(world_name, excel_name):
    doc = docx.Document(world_name)
    f = xlwt.Workbook()


    for i, table in enumerate(doc.tables):  # 遍历所有表格
       
        sheet = f.add_sheet('EAST' + str(i),cell_overwrite_ok=True) #创建sheet
        w = 0
        for row in table.rows:  # 遍历表格的所有行
            row_str = '|'.join([cell.text for cell in row.cells])  # 一行数据
            # print(row_str)
            li = row_str.split('|')
            li2 = list(set(li))
            li2.sort(key = li.index)
            if '传输文件名称' in li2:
                continue
            # print(li2)
        
            for j, k in enumerate(li2):
                sheet.write(w, j, k)
            w += 1

    f.save(excel_name)

# 读从world文档整理出来的excel
def read_excel(excel_name):
    li =[]
    li2 = []
    tables= []
    wb = xlrd.open_workbook(excel_name)

    for table in wb.sheets()[2:60]:
        nrows = table.nrows #行数
        ncols = table.ncols #列数
        # print(nrows, ncols)
        colnames =  table.row_values(1) #某一行数据
        # print(colnames)
        r = table.row_values(0)
        for rownum in range(0,nrows):
            row = table.row_values(rownum)
            if row and rownum >= 2:
                app = {}
                for i in range(len(colnames)):
                    app['传输文件名称'] = r[0]
                    app['表名'] = r[1]
                    app[colnames[i]] = row[i]
                    
                li.append(app)

    for table in wb.sheets()[60:]:
        nrows = table.nrows #行数
        ncols = table.ncols #列数
        # print(nrows, ncols)
        colnames =  table.row_values(0) #某一行数据 
        
        for rownum in range(1,nrows):
            row = table.row_values(rownum)
            # print(row)
            if row:
                app = {}
                for i in range(len(colnames)):
                    app[colnames[i]] = row[i]
            li2.append(app)
    # print(li)

    for i in li:
        i['格式'] = "DATE('yyyy-MM-dd')"
        # print(i)
        if i['数据元编码'] == '001005':

            i['格式'] = "DATE('yyyy-MM-dd')"
            # print(i['格式'])

        elif i['数据元编码'] == '001007' or i['数据元编码'] == '001008':
            i['格式'] = "DATE"

        else:
            for j in li2:

                if i['数据元编码'] == j ['数据元编码']:
                    i['格式'] = j['格式']

            
    return li, li2


# 获取表名
def get_table_name(file_name):
    sheet_list = []
    data, data2 = read_excel(file_name)
    for li in data:
        sheet_list.append((li['传输文件名称'],li['表名']))
    li = list(set(sheet_list))
    li.sort(key = sheet_list.index)
    # print(li)
    return li, data, data2

# 生成excel
def write_excel(file_name1, file_name2):
    choice = select_choice('DBDDEASTVFinal3.0.xlsx')
    f = xlwt.Workbook()
    row3 = ['Field name','Description','Field type (Format)','Mandatory','Chinese description','Filed value example','Remark','Is Search Field', 'UI input type']
    row4 = ['id','记录序列号','INT','M','Unique key for each entry','','','','']
    tables, data, data2 = get_table_name(file_name1)
    for w,table in enumerate(tables):
        sheet = f.add_sheet('EAST' + table[0],cell_overwrite_ok=True)
        sheet.write(0,0,'EAST'+ table[0] + '-east_' + table[0].lower() + ' ' + table[1])
        for i in range(0,len(row3)):
            # 第三行
            sheet.write(2,i,row3[i])  
            # 第四行
            sheet.write(3,i,row4[i])
        i = 0
        for j, li in enumerate(data):   
            if li['传输文件名称'] == table[0]:

                for m in choice:
                    # print('m---------------------',m)
                    # print('li-------------',li)
                    if m['table'][4:].replace(' ', '') == li['传输文件名称'] and m['Field name'] == li['数据项代码']:
                        req = m['Mandatory']
                        # if m['table'] == 'EASTXDYWDBHT':
                        #     print(m)

                        #     if m['Field name'] == 'DBHTH':
                        #         print(m)
                        #         print(li['传输文件名称'])


                if li['数据项代码'] == 'CJRQ':
                    vc = "DATE('yyyy-MM-dd')"

                elif li['格式'][0].lower() == 'c':
                    ret =  re.sub("\D", "", li['格式'])
                    vc = 'varchar(%s)' % ret

                elif li['格式'][0].lower() == 'i':
                    ret =  re.sub("\D", "", li['格式'])
                    if ret:
                        vc = 'int(%s)' % ret
                    else:
                        vc = 'int'

                elif li['格式'][0].lower() == 'f':
                    vc = 'varchar(10)'
                elif li['格式'] == "DATE('yyyy-MM-dd')":
                    vc = "DATE('yyyy-MM-dd')"

                elif li['格式'] == "DATE":
                    vc = "DATE"

                elif li['格式'][0].lower() == 'd':
                    # print(li['格式'])
                    ret = li['格式'][1:]
                    k = ret.split('.')
                    try:
                        vc = 'decimal(%s,%s)' % (k[0],k[1])
                    except:
                        vc = 'decimal(%s)' % (k[0])
                else:
                    vc = ''

                # print(i, table[0], li)
                sheet.write(i+4,0, li['数据项代码'])
                sheet.write(i+4,2, vc)
                sheet.write(i+4,3, req)
                sheet.write(i+4,4, li['数据项名称'])
                sheet.write(i+4,6, li['备注'])
                i += 1
            else:
                i = 0
    # print(data)           
    f.save(file_name2)

    rb = xlrd.open_workbook(file_name2, formatting_info=True)
    wb = copy(rb)
    tables = rb.sheets()
    i = 0
    for table in tables:
        
        nrows = table.nrows #行数
        # print(table.name,nrows)
        ws = wb.get_sheet(i)
        ws.write(nrows,0,'Audit fields')
        ws.write(nrows+1,0,'ediDate')
        ws.write(nrows+1,2,"DATE('yyyy-MM-dd')")
        ws.write(nrows+1,4,'数据日期')
        ws.write(nrows+2,0,'etlTimeStamp')
        ws.write(nrows+2,2,"TIMESTAMP('yyyy-MM-dd HH:MM:SS')")
        ws.write(nrows+2,4,'数据时间')
        ws.write(nrows+2,5,'1970-01-01 00:00:00')
        i += 1
    wb.save(file_name2)  



if __name__ == '__main__':
    world_2_excel(u'接口文档.docx', 'world_excel.xlsx')
    write_excel('world_excel.xlsx', '1234.xlsx')



print('excel生成完成:',time.time())

